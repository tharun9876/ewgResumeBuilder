from main import extract_docx_text, map_extraction_to_builder
from docx import Document
import io


def build_docx() -> bytes:
    doc = Document()
    doc.add_heading("Example Candidate", 0)
    doc.add_paragraph("Senior Oracle Utilities Developer")
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Acme | Senior Developer | Jan 2022 – Present")
    doc.add_paragraph("Developed batch processes and integrations.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Certifications"
    table.rows[0].cells[1].text = "Oracle Utilities CCS Specialist"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_docx_extraction_includes_paragraphs_and_tables():
    content = extract_docx_text(build_docx())
    assert "Example Candidate" in content
    assert "Developed batch processes" in content
    assert "Oracle Utilities CCS Specialist" in content


def test_competency_grid_keeps_each_table_paragraph_as_a_separate_point():
    import main

    doc = Document()
    doc.add_paragraph("Example Candidate")
    doc.add_paragraph("Key Skills")
    table = doc.add_table(rows=1, cols=3)
    groups = [
        ["Management Skills", "Project Planning & Execution", "People Management", "Customer Support"],
        ["Leadership Skills", "Facilitation & Negotiation", "Team Leadership & Collaboration"],
        ["Interpersonal Skills", "Quick Learner", "Good Communication Skills"],
    ]
    for cell, values in zip(table.rows[0].cells, groups):
        cell.paragraphs[0].text = values[0]
        for value in values[1:]:
            cell.add_paragraph(value)
    doc.add_paragraph("Professional Experience")
    doc.add_paragraph("QA Manager at Acme | Jan 2020 - Present")

    buffer = io.BytesIO()
    doc.save(buffer)
    source_text = extract_docx_text(buffer.getvalue())
    extracted, _, _ = main._extract_deterministic_resume(source_text)

    competency_texts = [block["text"] for block in extracted["competencies"]]
    assert competency_texts == [value for group in groups for value in group]
    assert all(block["type"] == "bullet" for block in extracted["competencies"])


def test_manual_line_breaks_stay_inside_their_source_responsibility():
    from docx.enum.text import WD_BREAK
    import main

    doc = Document()
    doc.add_paragraph("Example Candidate")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Acme Corp Jan 2020 - Present")
    first = doc.add_paragraph("• First responsibility line")
    first.add_run().add_break(WD_BREAK.LINE)
    first.add_run("continued on the next visual line")
    doc.add_paragraph("Continuation stored as another Word paragraph.")
    doc.add_paragraph("• Second responsibility")

    buffer = io.BytesIO()
    doc.save(buffer)
    source_text = extract_docx_text(buffer.getvalue())
    extracted, _, _ = main._extract_deterministic_resume(source_text)

    responsibilities = extracted["experience"][0]["responsibilities"]
    assert [item["text"] for item in responsibilities] == [
        "First responsibility line continued on the next visual line Continuation stored as another Word paragraph.",
        "Second responsibility",
    ]


def test_unheaded_summary_name_suffix_and_dated_employers_are_mapped_from_source():
    import main

    doc = Document()
    doc.add_paragraph("Jonathan", style="Heading 1")
    doc.add_paragraph("Ocampo US Citizen - Open to travel")
    source_summary = (
        "Certified Oracle Utilities consultant with more than twelve years of experience "
        "delivering implementation, upgrade, support, design, development, testing, "
        "requirements, business analysis, and team leadership across customer-focused "
        "projects. Recognized for strong problem solving and debugging skills."
    )
    doc.add_paragraph(source_summary)
    doc.add_paragraph("Experience", style="Heading 2")
    doc.add_paragraph("Orlando Utilities Commission 05/2026 - Present")
    doc.add_paragraph("Key Projects:")
    doc.add_paragraph("CC&B production support.")
    doc.add_paragraph("Wipro Limited 07/2017 - 02/2018")
    doc.add_paragraph("Key Projects:")
    doc.add_paragraph("Role Description:")
    doc.add_paragraph("Water billing customizations.")

    buffer = io.BytesIO()
    doc.save(buffer)
    source_text = extract_docx_text(buffer.getvalue())
    extracted, summary, _ = main._extract_deterministic_resume(source_text)
    notes = main._build_review_notes(extracted, source_text)

    assert extracted["full_name"] == "Jonathan Ocampo"
    assert not extracted["headline"]
    assert extracted["contact_details"] == "US Citizen - Open to travel"
    assert summary == source_summary
    assert extracted["professional_summary"] == source_summary
    assert [job["company"] for job in extracted["experience"]] == [
        "Orlando Utilities Commission", "Wipro Limited",
    ]
    assert all(not job["title"] for job in extracted["experience"])
    assert any("employer-level job title was not identified" in note for note in notes)
    assert all(
        "Role Description:" not in block["text"]
        for job in extracted["experience"]
        for block in job["responsibilities"]
    )
    assert not any("No Professional Summary" in note for note in notes)
    assert not any("unconfirmed employer/company" in note for note in notes)
    assert any("No email, phone number, or profile link" in note for note in notes)


def test_docx_package_rejects_non_docx_zip():
    import zipfile
    import pytest
    from fastapi import HTTPException
    from main import _validate_docx_package

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("not-a-docx.txt", "invalid")
    with pytest.raises(HTTPException) as error:
        _validate_docx_package(buffer.getvalue())
    assert error.value.status_code == 400


def test_docx_package_rejects_excessive_expanded_content(monkeypatch):
    import zipfile
    import pytest
    from fastapi import HTTPException
    import main

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "document")
        archive.writestr("word/media/image.bin", "x" * 100)
    monkeypatch.setattr(main, "MAX_DOCX_EXPANDED_BYTES", 50)
    with pytest.raises(HTTPException) as error:
        main._validate_docx_package(buffer.getvalue())
    assert error.value.status_code == 413


def test_upload_reader_stops_at_configured_limit(monkeypatch):
    import asyncio
    import pytest
    from fastapi import HTTPException
    import main

    class FakeUpload:
        async def read(self, size):
            assert size == 6
            return b"123456"

    monkeypatch.setattr(main, "MAX_UPLOAD_BYTES", 5)
    with pytest.raises(HTTPException) as error:
        asyncio.run(main._read_upload_limited(FakeUpload()))
    assert error.value.status_code == 413


def test_live_sections_are_built_from_partial_extraction():
    from main import _live_sections_from_extraction
    extracted = _expected_extraction("Kiran Example")
    extracted["contact_details"] = "kiran@example.com"
    extracted["competencies"] = [{"type": "bullet", "text": "Oracle Utilities / OUAF", "style": "disc"}]
    extracted["experience"] = [{
        "company": "Example Utilities", "title": "Developer", "start_date": "Jan 2023",
        "end_date": "Present", "client_or_project": "CCS Support", "project_context": "Supported Meter-to-Cash.",
        "responsibilities": [{"type": "bullet", "text": "Built batch processes.", "style": "disc"}],
    }]
    sections = _live_sections_from_extraction(extracted)
    by_id = {section["id"]: section for section in sections}
    assert by_id["profile"]["text"] == "Kiran Example\nOracle Utilities SME\nkiran@example.com"
    assert "Oracle Utilities / OUAF" in by_id["competencies"]["text"]
    assert "Example Utilities — Developer" in by_id["experience"]["text"]
    assert "Built batch processes." in by_id["experience"]["text"]
    assert "review_notes" not in by_id


def test_section_aware_chunker_keeps_source_sections_separate():
    from main import _split_resume_into_section_chunks
    text = """Venkata Kiran Omkaram
Oracle Utilities SME
Professional Summary
Experienced in Oracle Utilities.
Certifications
Oracle Utilities CCS Specialist
Key Expertise
Upgrade management
Work Experience
Developer at Example Corp | Jan 2022 – Present
Built batch processes.
"""
    chunks = _split_resume_into_section_chunks(text, max_chars=1600, overlap=150)
    hints = [hint for _, hint in chunks]
    assert hints == ["profile", "summary", "certifications", "competencies", "experience"]
    assert all("Built batch processes" not in chunk for chunk, hint in chunks if hint != "experience")


def test_source_coverage_audit_preserves_unmapped_section_lines_for_review():
    import main

    extracted = main._empty_extraction()
    extracted["additional_sections"] = []
    groups = {"certifications": ["Oracle Utilities Certified Implementation Specialist"]}
    gaps = main._preserve_unmapped_section_content(extracted, groups, {"certifications"})
    notes = main._build_review_notes(extracted, "Certifications\nOracle Utilities Certified Implementation Specialist")

    assert gaps == [{"section": "Certifications", "items": ["Oracle Utilities Certified Implementation Specialist"]}]
    preserved = extracted["additional_sections"][-1]
    assert preserved["heading"] == "Source Content to Review - Certifications"
    assert preserved["items"][0]["text"] == "Oracle Utilities Certified Implementation Specialist"
    assert any("preserved in 'Source Content to Review - Certifications'" in note for note in notes)


def test_section_coverage_audit_does_not_duplicate_content_already_mapped():
    import main

    extracted = main._empty_extraction()
    extracted["certifications"] = ["Oracle Utilities Certified Implementation Specialist"]
    extracted["additional_sections"] = []
    gaps = main._preserve_unmapped_section_content(
        extracted,
        {"certifications": ["Oracle Utilities Certified Implementation Specialist"]},
        {"certifications"},
    )

    assert gaps == []
    assert extracted["additional_sections"] == []


def test_common_heading_variants_are_recognized_for_section_mapping():
    from main import _resume_section_heading

    assert _resume_section_heading("Technical Knowledge") == "skills"
    assert _resume_section_heading("Education and Training") == "education"
    assert _resume_section_heading("Professional Licenses") == "certifications"
    assert _resume_section_heading("Selected Accomplishments") == "achievements"


def test_full_document_audit_keeps_content_from_unknown_or_missing_sections():
    import main

    extracted = main._empty_extraction()
    extracted.update({
        "full_name": "Jane Candidate",
        "additional_sections": [],
        "source_text_for_coverage": "Jane Candidate\nCommunity Leadership\nVolunteer coordinator for local food bank.",
    })
    gaps = main._preserve_unmapped_section_content(extracted, {"profile": ["Jane Candidate"]}, set())

    review_section = extracted["additional_sections"][-1]
    assert review_section["heading"] == "Source Content to Review - Unmapped Content"
    assert [item["text"] for item in review_section["items"]] == [
        "Community Leadership", "Volunteer coordinator for local food bank.",
    ]
    assert gaps[-1]["section"] == "Unmapped Content"


def test_extraction_is_restricted_to_detected_section():
    from main import _restrict_chunk_to_section
    raw = _expected_extraction("Kiran Example")
    raw["competencies"] = [{"type": "bullet", "text": "Upgrade management", "style": "disc"}]
    raw["experience"] = [{"title": "Developer", "responsibilities": [{"type": "bullet", "text": "Built batches", "style": "disc"}]}]
    raw["review_notes"] = ["No education found in this chunk."]
    filtered = _restrict_chunk_to_section(raw, "competencies")
    assert filtered["competencies"] == raw["competencies"]
    assert filtered["experience"] == []
    assert filtered["professional_summary"] == ""
    assert filtered["review_notes"] == []


def test_section_json_schema_only_requests_relevant_fields():
    from main import _schema_for_section
    experience_schema = _schema_for_section("experience")
    assert experience_schema["required"] == ["experience"]
    assert set(experience_schema["properties"]) == {"experience"}
    general_schema = _schema_for_section("general")
    assert "full_name" in general_schema["properties"]
    assert "experience" in general_schema["properties"]


def test_review_notes_are_derived_from_complete_resume_not_chunk_boilerplate():
    from main import _build_review_notes
    text = "Kiran Example\nProfessional Summary\nExperienced engineer.\nWork Experience\nDeveloper | Jan 2022 – Present\nBuilt APIs."
    extracted = _expected_extraction("Kiran Example")
    extracted["headline"] = "Engineer"
    extracted["review_notes"] = ["No full name provided in the chunk."]
    notes = _build_review_notes(extracted, text)
    assert any("No Education section" in note for note in notes)
    assert not any("No full name provided in the chunk" in note for note in notes)


def test_mapping_matches_builder_shape():
    extracted = {
        "full_name": "Example Candidate",
        "headline": "Oracle Utilities SME",
        "contact_details": "test@example.com",
        "professional_summary": "Experienced professional.",
        "competencies": [{"type": "bullet", "text": "OUAF", "style": "disc"}],
        "education": [],
        "experience": [{"company": "Acme", "title": "Developer", "start_date": "Jan 2022", "end_date": "Present", "client_or_project": "Utility project", "project_context": "Implemented a new integration.", "responsibilities": [{"type": "bullet", "text": "Built batches.", "style": "disc"}]}],
        "technical_skills": [{"type": "heading", "text": "Oracle Utilities", "style": "disc"}],
        "certifications": ["Oracle Utilities CCS Specialist"],
        "achievements": [],
        "additional_sections": [{"heading": "Languages", "items": [{"type": "bullet", "text": "Java", "style": "disc"}]}],
        "review_notes": ["Education was not found."],
    }
    data, notes = map_extraction_to_builder(extracted)
    assert data["name"] == "Example Candidate"
    assert data["contactDetails"] == "test@example.com"
    assert data["experience"][0]["company"] == "Acme"
    assert data["experience"][0]["blocks"][0]["type"] == "paragraph"
    assert "certs" in data["sectionOrder"]
    assert any(key.startswith("custom_") for key in data["customSections"])
    assert notes == ["Education was not found."]


def _expected_extraction(full_name="Example Candidate"):
    return {
        "full_name": full_name,
        "headline": "Oracle Utilities SME",
        "contact_details": "",
        "professional_summary": "Experienced professional.",
        "competencies": [],
        "education": [],
        "experience": [],
        "technical_skills": [],
        "certifications": [],
        "achievements": [],
        "additional_sections": [],
        "review_notes": [],
    }


def _fake_stream_response(content, finish_reason="stop", status_code=200):
    import json

    class FakeResponse:
        def __init__(self):
            self.status_code = status_code
            self.lines = [
                "data: " + json.dumps({"choices": [{"delta": {"content": content}}]}),
                "data: " + json.dumps({"choices": [{"delta": {}, "finish_reason": finish_reason}]}),
                "data: [DONE]",
            ]

        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        async def aiter_lines(self):
            for line in self.lines:
                yield line
        async def aread(self): return b"provider error"

    return FakeResponse()


def test_openrouter_extraction_uses_streaming_summary_refinement(monkeypatch):
    import asyncio
    import json
    import httpx
    from main import analyse_with_openrouter

    source = """Kiran Example
Oracle Utilities SME
Professional Summary
Experienced Oracle Utilities professional who builds reliable business services and batch programs.
Certifications
Oracle Utilities CCS Specialist
Key Expertise
OUAF development
Project delivery
Work Experience
Developer at Acme Utilities | Jan 2022 – Present
Built batch processes and integrations.
"""
    model_response = {"professional_summary": "Experienced Oracle Utilities professional who builds reliable business services and batch programs."}

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs): self.timeout = kwargs.get("timeout")
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        def stream(self, method, url, headers, json):
            assert method == "POST"
            assert url == "https://openrouter.ai/api/v1/chat/completions"
            assert headers["Authorization"] == "Bearer test-key"
            assert json["model"] == "openrouter/free"
            assert json["response_format"]["type"] == "json_schema"
            assert json["response_format"]["json_schema"]["name"] == "resume_extraction_summary"
            assert set(json["response_format"]["json_schema"]["schema"]["properties"]) == {"professional_summary"}
            assert json["response_format"]["json_schema"]["strict"] is True
            assert json["provider"]["require_parameters"] is True
            assert json["stream"] is True
            assert "tools" not in json
            return _fake_stream_response(__import__('json').dumps(model_response))

    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setenv("OPENROUTER_MODEL", "openrouter/free")
    monkeypatch.setattr(httpx, "AsyncClient", FakeAsyncClient)
    streamed_events = []

    async def progress(event): streamed_events.append(event)
    result = asyncio.run(analyse_with_openrouter(source, progress_callback=progress))
    assert result["full_name"] == "Kiran Example"
    assert result["headline"] == "Oracle Utilities SME"
    assert result["professional_summary"] == model_response["professional_summary"]
    assert [x["text"] for x in result["competencies"]] == ["OUAF development", "Project delivery"]
    assert result["certifications"] == ["Oracle Utilities CCS Specialist"]
    assert result["experience"][0]["company"] == "Acme Utilities"
    assert result["experience"][0]["title"] == "Developer"
    assert any(e["type"] == "request_started" for e in streamed_events)
    assert any(e["type"] == "request_finished" for e in streamed_events)
    assert any(e["type"] == "chunk_completed" for e in streamed_events)
    assert any(e["type"] == "section_update" for e in streamed_events)
    section_event = next(e for e in streamed_events if e["type"] == "section_update")
    assert any(section["id"] == "experience" for section in section_event["sections"])


def test_single_competency_paragraph_is_split_only_when_source_words_are_preserved(monkeypatch):
    import asyncio
    import json
    import httpx
    import main

    source = """Example Candidate
Professional Summary
Experienced QA manager with utility experience.
Key Expertise
Management Skills Project Planning & Execution People Management Customer Support
Work Experience
QA Manager at Example Utilities | Jan 2022 - Present
Led user acceptance testing.
"""

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        def stream(self, method, url, headers, json):
            schema_name = json["response_format"]["json_schema"]["name"]
            if schema_name == "resume_extraction_summary":
                result = {"professional_summary": "QA manager with utility experience."}
            else:
                assert schema_name == "resume_extraction_competencies"
                result = {"competencies": [
                    {"type": "bullet", "text": "Management Skills", "style": "disc"},
                    {"type": "bullet", "text": "Project Planning & Execution", "style": "disc"},
                    {"type": "bullet", "text": "People Management", "style": "disc"},
                    {"type": "bullet", "text": "Customer Support", "style": "disc"},
                ]}
            return _fake_stream_response(json_module.dumps(result))

    json_module = json
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setenv("OPENROUTER_MODEL", "openrouter/free")
    monkeypatch.setattr(httpx, "AsyncClient", FakeAsyncClient)
    result = asyncio.run(main.analyse_with_openrouter(source))
    assert [item["text"] for item in result["competencies"]] == [
        "Management Skills", "Project Planning & Execution", "People Management", "Customer Support",
    ]


def test_deterministic_parser_preserves_actual_source_sections_and_roles():
    from main import _extract_deterministic_resume
    source = """Example Candidate
Oracle Utilities SME
Professional Summary
Implemented utility billing solutions.
Worked on batch performance and integrations.
Certifications
Oracle Utilities CCS Specialist
Key Expertise
Billing and meter operations
OUAF batch development
Work Experience
Senior Developer at Alpha Ltd | Jan 2023 – Present
Implemented batch processing.
Developer | May 2020 – Dec 2022
Supported integrations for a utility.
• Built SOAP services.
"""
    extracted, summary_source, headings = _extract_deterministic_resume(source)
    assert extracted["full_name"] == "Example Candidate"
    assert extracted["headline"] == "Oracle Utilities SME"
    assert "billing solutions" in summary_source
    assert [item["text"] for item in extracted["competencies"]] == ["Billing and meter operations", "OUAF batch development"]
    assert extracted["certifications"] == ["Oracle Utilities CCS Specialist"]
    assert len(extracted["experience"]) == 2
    assert extracted["experience"][0]["company"] == "Alpha Ltd"
    assert extracted["experience"][0]["title"] == "Senior Developer"
    assert extracted["experience"][1]["title"] == "Developer"
    assert [item["text"] for item in extracted["experience"][1]["responsibilities"]] == [
        "Supported integrations for a utility.", "Built SOAP services."
    ]
    assert extracted["education"] == []
    assert extracted["achievements"] == []
    assert headings == {"summary", "certifications", "competencies", "experience"}


def test_deterministic_parser_preserves_education_skills_and_objective_heading():
    from main import _extract_deterministic_resume, map_extraction_to_builder
    source = """Example Candidate
Oracle Utilities Consultant
Objective
Experienced consultant with utility billing expertise.
Education
KANSAS STATE UNIVERSITY, Manhattan, Kansas
Bachelor of Science in Agriculture
Emphasis in Economics and Marketing
HUTCHINSON COMMUNITY COLLEGE, Hutchinson, Kansas
Associate of Arts in Agricultural Economics
Technical Skills
Oracle CC&B, C2M, OUAF
Applications: Oracle Utilities SDK, OIC, IWS
Key Expertise
Requirements management
Work Experience
Consultant at Example Corp | Jan 2022 - Present
Led delivery activities.
"""
    extracted, _, _ = _extract_deterministic_resume(source)
    data, _ = map_extraction_to_builder(extracted)
    assert len(extracted["education"]) == 2
    assert extracted["education"][0] == {
        "degree": "Bachelor of Science in Agriculture; Emphasis in Economics and Marketing",
        "institution": "KANSAS STATE UNIVERSITY, Manhattan, Kansas", "year": "",
    }
    assert [block["text"] for block in extracted["technical_skills"]] == [
        "Oracle CC&B, C2M, OUAF", "Applications", "Oracle Utilities SDK", "OIC", "IWS",
    ]
    assert data["headings"]["summary"] == "OBJECTIVE"
    assert [item["text"] for item in extracted["competencies"]] == ["Requirements management"]


def test_hierarchical_experience_keeps_clients_and_parent_employer_context_separate():
    from main import _parse_source_experience
    roles = _parse_source_experience([
        "THE MARKETING GROUP (TMG), Sarasota, Florida | January 2022 - November 2024",
        "Project Manager | January 2022 - November 2024",
        "Client: SWBNO | January 2022 - November 2024",
        "Managed implementation planning.",
        "Business Analyst | January 2021 - December 2021",
        "Documented requirements.",
    ])
    assert len(roles) == 2
    assert all(role["company"].startswith("THE MARKETING GROUP") for role in roles)
    assert roles[0]["client_or_project"] == "SWBNO"
    assert roles[0]["responsibilities"][0]["text"] == "Managed implementation planning."
    assert "Employer tenure:" in roles[0]["project_context"]
    assert roles[1]["responsibilities"][0]["text"] == "Documented requirements."


def test_experience_responsibilities_do_not_leak_into_competencies():
    from main import _extract_deterministic_resume
    source = """Example Candidate
Objective
Utility professional.
Core Competencies
Customer engagement
Experience
Analyst at Example Corp | Jan 2022 - Present
Led customer engagement for billing modernization.
Coordinated testing with business teams.
"""
    extracted, _, _ = _extract_deterministic_resume(source)
    assert [item["text"] for item in extracted["competencies"]] == ["Customer engagement"]
    assert [item["text"] for item in extracted["experience"][0]["responsibilities"]] == [
        "Led customer engagement for billing modernization.",
        "Coordinated testing with business teams.",
    ]


def test_source_lines_with_separate_title_company_dates_and_employer_are_associated():
    from main import _parse_source_experience
    roles = _parse_source_experience([
        "Sr. QA Manager",
        "Berkshire Hathaway Energy, Portland, OR",
        "Jan 2016 - Present",
        "Projects Handled: Customer Transformation, GTx Replacement",
        "Technologies: Oracle C2M, Maximo, ADO",
        "Responsibilities:",
        "Led end-to-end testing across integrated systems.",
        "• Developed unified test strategies.",
        "Test Project Manager",
        "The Standard Insurance Company, Portland, OR",
        "Feb 2010 - Jan 2016",
        "Employer: IBM",
        "Responsibilities:",
        "Managed offshore QA teams.",
    ])
    assert len(roles) == 2
    assert roles[0]["company"] == "Berkshire Hathaway Energy, Portland, OR"
    assert roles[0]["title"] == "Sr. QA Manager"
    assert (roles[0]["start_date"], roles[0]["end_date"]) == ("Jan 2016", "Present")
    assert roles[0]["client_or_project"] == "Customer Transformation, GTx Replacement"
    assert "Technologies: Oracle C2M, Maximo, ADO" in roles[0]["project_context"]
    assert [item["text"] for item in roles[0]["responsibilities"]] == [
        "Led end-to-end testing across integrated systems.", "Developed unified test strategies."
    ]
    assert roles[1]["company"] == "IBM"
    assert roles[1]["client_or_project"].startswith("The Standard Insurance Company")


def test_concise_summary_fallback_selects_source_statements_and_utility_fact():
    from main import _concise_source_summary, _extract_deterministic_resume, map_extraction_to_builder
    summary = _concise_source_summary([
        "Leading/Management",
        "Results-driven QA Manager with 20 years of experience leading testing projects.",
        "Proven expertise managing cross-functional teams and delivering projects on time.",
        "Highly skilled in Agile methodologies and risk management.",
        "Extensive experience leading enterprise testing, including IBM Maximo.",
        "Proven ability to define test strategies and ensure data integrity.",
        "Energy & Utilities",
        "10+ years of experience in the utility industry.",
        "Others",
        "Knowledge of Python, HTTP, HTML, JavaScript.",
    ])
    assert "10+ years of experience in the utility industry." in summary
    assert "Knowledge of Python" not in summary
    assert len(summary.split()) <= 110
    extracted, _, _ = _extract_deterministic_resume(
        "Example Candidate\nCareer Summary\n" + "\n".join([
            "Leading/Management", "Results-driven QA Manager with 20 years of experience leading testing projects.",
            "Proven expertise managing cross-functional teams and delivering projects on time.",
            "Highly skilled in Agile methodologies and risk management.",
            "Extensive experience leading enterprise testing, including IBM Maximo.",
            "Proven ability to define test strategies and ensure data integrity.",
            "Energy & Utilities", "10+ years of experience in the utility industry.",
            "Others", "Knowledge of Python, HTTP, HTML, JavaScript.",
        ])
    )
    assert len(extracted["professional_summary"].split()) <= 110
    assert "Knowledge of Python" in extracted["source_summary"]
    builder_data, _ = map_extraction_to_builder(extracted)
    assert "Knowledge of Python" in builder_data["sourceSummary"]


def test_ai_summary_grounding_rejects_new_claims_and_numbers():
    from main import _summary_is_source_grounded

    source = "QA Manager with 20 years of experience leading testing programs in utilities."
    assert _summary_is_source_grounded("QA Manager with 20 years experience in utilities.", source)
    assert not _summary_is_source_grounded("QA Manager with 25 years experience in utilities.", source)
    assert not _summary_is_source_grounded("Award-winning QA Manager with 20 years experience in utilities.", source)


def test_review_panel_is_outside_scrolling_form_and_toggleable():
    from pathlib import Path
    html = (Path(__file__).parent / "index.html").read_text(encoding="utf-8")
    review_pos = html.index('<div id="resumeReviewNotes"')
    steps_pos = html.index('<div class="steps-track" id="stepsTrack">')
    editor_body_pos = html.index('<div class="editor-body" id="editorBody">')
    assert review_pos < steps_pos < editor_body_pos
    assert "function toggleResumeReviewNotes()" in html
    assert "Review these AI extraction notes before exporting:" in html
    assert "function focusResumeReviewNotes()" in html
    assert "Review AI extraction notes" in html
    assert "const displayName = hasCompany ? j.company : j.role;" in html
    assert 'id="paneResizer" role="separator"' in html
    assert "setPointerCapture(event.pointerId)" in html
    assert "ewgResumeBuilderEditorWidth" in html
    assert "paneResizer.addEventListener('keydown'" in html
    assert "MAX_RESUME_RESPONSIBILITIES_PER_ROLE = 4" in html
    assert "Include all extracted responsibilities in the resume" in html
    assert "View full source career summary (reference only)" in html
    assert 'id="resumeChatPanel"' in html
    assert "fetch('/api/refine-resume'" in html
    assert "function applyResumeChatChanges(changes)" in html


def test_chunked_extraction_splits_when_model_finishes_at_length(monkeypatch):
    import asyncio
    import json
    import httpx
    from main import _request_json_chunk

    left_data = _expected_extraction("Chunk Candidate")
    right_data = _expected_extraction("")
    left_data["experience"] = [{
        "company": "Acme", "title": "Developer", "start_date": "Jan 2022", "end_date": "Present",
        "client_or_project": "", "project_context": "", "responsibilities": [
            {"type": "bullet", "text": "Built a batch process.", "style": "disc"}
        ]
    }]
    right_data["experience"] = [{
        "company": "Acme", "title": "Developer", "start_date": "Jan 2022", "end_date": "Present",
        "client_or_project": "", "project_context": "", "responsibilities": [
            {"type": "bullet", "text": "Built a batch process.", "style": "disc"},
            {"type": "bullet", "text": "Maintained integrations.", "style": "disc"}
        ]
    }]

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs): self.calls = 0
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        def stream(self, method, url, headers, json):
            self.calls += 1
            if self.calls == 1:
                return _fake_stream_response('{"partial":', finish_reason="length")
            payload = left_data if self.calls == 2 else right_data
            return _fake_stream_response(__import__('json').dumps(payload))

    monkeypatch.setenv("OPENROUTER_CHUNK_MAX_TOKENS", "4096")
    async def run():
        client = FakeAsyncClient()
        async with client:
            return await _request_json_chunk(
                client=client,
                endpoint="https://openrouter.ai/api/v1/chat/completions",
                headers={"Authorization": "Bearer test-key"},
                model="openrouter/free",
                chunk="This is a long resume passage. " * 100,
                chunk_number=1,
                chunk_count=1,
                known_roles=[],
                budget={"calls": 0},
            )
    result = asyncio.run(run())
    assert len(result["experience"]) == 1
    bullets = [item["text"] for item in result["experience"][0]["responsibilities"]]
    assert bullets == ["Built a batch process.", "Maintained integrations."]


def test_split_text_chunks_are_bounded_and_cover_content():
    from main import _split_text_into_chunks
    content = "\n".join(f"Resume line {i} with useful details." for i in range(300))
    chunks = _split_text_into_chunks(content, max_chars=1800, overlap=200)
    assert len(chunks) > 1
    assert all(len(chunk) <= 1800 for chunk in chunks)
    assert "Resume line 0" in chunks[0]
    assert "Resume line 299" in chunks[-1]

def test_stream_endpoint_sends_sse_progress_and_complete_payload(monkeypatch):
    import asyncio
    from fastapi.testclient import TestClient
    import main

    monkeypatch.setattr(main, "extract_docx_text", lambda _content: "Example resume text")

    async def fake_analyse(text, progress_callback=None):
        assert text == "Example resume text"
        if progress_callback:
            await progress_callback({"type": "analysis_started", "text_chars": len(text), "total_chunks": 1, "model": "mock/free"})
            await progress_callback({"type": "token_progress", "chunk_number": 1, "chunk_count": 1, "depth": 0, "api_call": 1, "chars_received": 200})
            await progress_callback({"type": "section_update", "chunk_number": 1, "chunk_count": 1, "sections": [{"id": "summary", "title": "Professional Summary", "text": "Experienced professional."}]})
            await progress_callback({"type": "chunk_completed", "chunk_number": 1, "chunk_count": 1, "chunks_completed": 1, "total_chunks": 1})
        return _expected_extraction()

    monkeypatch.setattr(main, "analyse_with_openrouter", fake_analyse)
    with TestClient(main.app) as client:
        response = client.post(
            "/api/analyze-resume-stream",
            files={"file": ("sample.docx", b"fake docx bytes", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/event-stream")
    assert '"type":"analysis_started"' in response.text
    assert '"type":"token_progress"' in response.text
    assert '"type":"section_update"' in response.text
    assert '"title":"Professional Summary"' in response.text
    assert '"type":"complete"' in response.text
    assert '"source_filename":"sample.docx"' in response.text
    assert response.text.endswith("data: [DONE]\n\n")



def test_summary_refinement_failure_keeps_source_content_and_review_note(monkeypatch):
    import asyncio
    import httpx
    import main

    source = """Example Candidate
Oracle Utilities Developer
Professional Summary
Experienced Oracle Utilities developer working on billing, integrations, and batch processing.
Certifications
Oracle Utilities CCS Specialist
Key Expertise
OUAF development
Work Experience
Developer at Acme Utilities | Jan 2022 – Present
Built utility billing integrations.
"""

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs): self.timeout = kwargs.get("timeout")
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False

    async def failing_request(**kwargs):
        raise RuntimeError("simulated free-model truncation")

    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setenv("OPENROUTER_MODEL", "openrouter/free")
    monkeypatch.setattr(httpx, "AsyncClient", FakeAsyncClient)
    monkeypatch.setattr(main, "_request_json_chunk", failing_request)

    result = asyncio.run(main.analyse_with_openrouter(source))
    assert "Experienced Oracle Utilities developer" in result["professional_summary"]
    assert any("concise source-based summary was retained" in note for note in result["review_notes"])


def test_resume_chat_returns_section_scoped_edits_without_external_request(monkeypatch):
    import asyncio
    import json
    import httpx
    import main

    class FakeResponse:
        status_code = 200
        def json(self):
            return {"choices": [{"message": {"content": json.dumps({
                "reply": "I shortened the summary using the supplied resume facts.",
                "changes": [
                    {"section": "summary", "value": "QA leader with 20 years of experience."},
                    {"section": "unknown", "value": "should be ignored"},
                ],
            })}}]}

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        async def post(self, url, headers, json):
            assert url == "https://openrouter.ai/api/v1/chat/completions"
            assert json["response_format"]["json_schema"]["name"] == "resume_chat_edit"
            assert "Never invent or infer" in json["messages"][0]["content"]
            assert "sourceSummary" in json["messages"][1]["content"]
            return FakeResponse()

    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setenv("OPENROUTER_MODEL", "openrouter/free")
    monkeypatch.setattr(httpx, "AsyncClient", FakeAsyncClient)
    result = asyncio.run(main.refine_resume({
        "instruction": "Shorten the career summary.",
        "resume": {"professional_summary": "Long original.", "sourceSummary": "QA leader with 20 years."},
        "history": [],
    }))
    assert result["reply"] == "Updated the requested section: Professional Summary."
    assert result["changes"] == [{"section": "summary", "value": "QA leader with 20 years of experience."}]


def test_headline_request_cannot_report_success_from_a_summary_only_change(monkeypatch):
    import asyncio
    import json
    import httpx
    import main

    class FakeResponse:
        status_code = 200
        def json(self):
            return {"choices": [{"message": {"content": json.dumps({
                "reply": "I've added the title/headline. Here's the updated summary field.",
                "changes": [{"section": "summary", "value": "A changed summary that was not requested."}],
            })}}]}

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        async def post(self, url, headers, json): return FakeResponse()

    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setattr(httpx, "AsyncClient", FakeAsyncClient)
    result = asyncio.run(main.refine_resume({
        "instruction": "add title/headline",
        "resume": {
            "headline": "", "professional_summary": "Experienced project manager.",
            "sourceSummary": "Experienced project manager in utilities.",
            "experience": [{"role": "Project Manager / Deployment Lead (IFS MWM Project)"}],
        },
        "history": [],
    }))

    assert result["changes"] == [{"section": "headline", "value": "Project Manager / Deployment Lead"}]
    assert result["reply"] == "Added the source-supported headline: Project Manager / Deployment Lead."


def test_resume_chat_scopes_followups_to_the_last_explicit_section():
    import main

    assert main._requested_resume_chat_sections("Shorten the professional summary") == {"summary"}
    assert main._requested_resume_chat_sections(
        "Make it more concise",
        [{"role": "user", "content": "Add separate points under Core Competences"}],
    ) == {"competencies"}
    assert main._requested_resume_chat_sections(
        "Shorten the Summary of Experience section",
    ) == {"custom_sections"}
    assert main._requested_resume_chat_sections("Move Education before Experience") == {"section_order"}


def test_resume_chat_detects_noop_field_edits():
    import json
    import main

    assert main._resume_chat_change_is_noop("summary", "Current summary", {"professional_summary": "Current summary"})
    assert main._resume_chat_change_is_noop("competencies", json.dumps([{"type": "bullet", "text": "Planning", "style": "disc"}]), {
        "competencies": [{"id": "local-only", "type": "bullet", "text": "Planning", "style": "disc"}],
    })
    assert not main._resume_chat_change_is_noop("summary", "New summary", {"professional_summary": "Current summary"})


def test_summary_request_discards_an_unrequested_headline_change(monkeypatch):
    import asyncio
    import json
    import httpx
    import main

    class FakeResponse:
        status_code = 200
        def json(self):
            return {"choices": [{"message": {"content": json.dumps({
                "reply": "Updated your headline.",
                "changes": [{"section": "headline", "value": "Unrequested Title"}],
            })}}]}

    class FakeAsyncClient:
        def __init__(self, *args, **kwargs): pass
        async def __aenter__(self): return self
        async def __aexit__(self, *args): return False
        async def post(self, url, headers, json): return FakeResponse()

    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setattr(httpx, "AsyncClient", FakeAsyncClient)
    result = asyncio.run(main.refine_resume({
        "instruction": "Shorten the professional summary",
        "resume": {"professional_summary": "Long summary", "headline": ""},
        "history": [],
    }))

    assert result["changes"] == []
    assert "Professional Summary" in result["reply"]
    assert "No resume fields were changed" in result["reply"]


def test_resume_chat_rejects_empty_instruction():
    import asyncio
    import pytest
    from fastapi import HTTPException
    from main import refine_resume
    with pytest.raises(HTTPException) as error:
        asyncio.run(refine_resume({"instruction": "  ", "resume": {}}))
    assert error.value.status_code == 400
